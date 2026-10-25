"""Spreadsheet export (XLSX, opens in Numbers or Excel) of the current snapshot plus history."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import analysis
from .snapshot import Snapshot
from .store import Store

FONT_NAME = "Arial"
FONT_SIZE = 12
BOLD = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)
REGULAR = Font(name=FONT_NAME, size=FONT_SIZE)


def _sanitize(value: Any) -> Any:
    if isinstance(value, str):
        return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    if isinstance(value, dict):
        return ", ".join(f"{k}={v}" for k, v in value.items())
    if isinstance(value, float) and value != value:
        return None
    return value


def _cell(value: Any) -> Any:
    return _sanitize(value)


PERCENT_KEYS = {"load", "lineUsage", "rating", "maintenanceState", "productionRating"}


def _format_cells(
    ws: Worksheet, percent_columns: set[int] = frozenset(), plain_columns: set[int] = frozenset()
) -> None:
    """Thousands separators for amounts, percent for ratio columns, none for ids (1-based column indexes)."""
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, bool) or not isinstance(cell.value, (int, float)):
                continue
            if cell.column in plain_columns:
                cell.number_format = "0"
            elif cell.column in percent_columns:
                cell.number_format = "0%"
            elif isinstance(cell.value, int):
                cell.number_format = "#,##0"
            else:
                cell.number_format = "#,##0.00"


def _table(ws: Worksheet, columns: list[tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    """columns: (header shown, key in row)."""
    ws.append([header for header, _ in columns])
    for cell in ws[1]:
        cell.font = BOLD
    for row in rows:
        ws.append([_cell(row.get(key)) for _, key in columns])
    ws.freeze_panes = "A2"
    for index, (header, key) in enumerate(columns, start=1):
        width = max([len(str(header))] + [len(str(_cell(r.get(key)) or "")) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 8), 60)
    _format_cells(
        ws,
        {i for i, (_, key) in enumerate(columns, start=1) if key in PERCENT_KEYS},
        {i for i, (_, key) in enumerate(columns, start=1) if key == "id"},
    )


def _summary(ws: Worksheet, snap: Snapshot, overview: dict[str, Any]) -> None:
    ws.append(["Snapshot", snap.date_label])
    prev = overview.get("previous") or {}
    rows = [
        ("Balance", overview.get("balance"), prev.get("balanceDelta")),
        ("Loan", overview.get("loan"), prev.get("loanDelta")),
        ("Earnings (calculateEarnings)", overview.get("earnings"), None),
        ("Total annual line profit", overview.get("linesProfitYear"), None),
        ("Lines", overview.get("lines"), prev.get("linesDelta")),
        ("Loss-making lines", overview.get("linesLosingMoney"), None),
        ("Vehicles", overview.get("vehicles"), prev.get("vehiclesDelta")),
        ("Passengers transported", overview.get("passengersTransported"), None),
        ("Cargo transported", overview.get("cargoTransported"), None),
        ("Collection errors", overview.get("snapshotErrors"), None),
    ]
    ws.append([])
    ws.append(["Indicator", "Value", f"Change since {prev.get('date', '-')}"])
    for cell in ws[3]:
        cell.font = BOLD
    for row in rows:
        ws.append(list(row))
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 26
    _format_cells(ws)


def _line_chart(ws: Worksheet, title: str, y_title: str, data_cols: list[int], rows: int, anchor: str) -> None:
    if rows < 2:
        return
    chart = LineChart()
    chart.title = title
    chart.y_axis.title = y_title
    chart.height, chart.width = 8, 18
    for col in data_cols:
        chart.add_data(Reference(ws, min_col=col, min_row=1, max_row=rows + 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=rows + 1))
    _show_axes(chart)
    ws.add_chart(chart, anchor)


CHART_WIDTH_CM = 26
CHART_HEIGHT_CM = 12
CHART_ROW_STEP = 30  # rows between chart anchors on the charts sheet (Arial 12 rows ≈ 0.6 cm)


def _bar_chart(
    canvas: Worksheet,
    data: Worksheet,
    title: str,
    label_col: int,
    value_col: int,
    rows: int,
    anchor: str,
    header_row: int = 1,
    horizontal: bool = True,
) -> None:
    """Bar chart on `canvas` over a table in `data` (header at `header_row`, then `rows` rows)."""
    if rows < 1:
        return
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.title = title
    chart.height, chart.width = CHART_HEIGHT_CM, CHART_WIDTH_CM
    chart.legend = None
    chart.add_data(
        Reference(data, min_col=value_col, min_row=header_row, max_row=header_row + rows), titles_from_data=True
    )
    chart.set_categories(Reference(data, min_col=label_col, min_row=header_row + 1, max_row=header_row + rows))
    _show_axes(chart)
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.dataLabels.numFmt = "#,##0"
    canvas.add_chart(chart, anchor)


def _show_axes(chart: BarChart | LineChart) -> None:
    """openpyxl marks axes as deleted by default; Numbers then hides scales and labels."""
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.y_axis.numFmt = "#,##0"


def _pie_chart(
    canvas: Worksheet,
    data: Worksheet,
    title: str,
    label_col: int,
    value_col: int,
    rows: int,
    anchor: str,
    header_row: int,
) -> None:
    if rows < 1:
        return
    chart = PieChart()
    chart.title = title
    chart.height, chart.width = CHART_HEIGHT_CM, CHART_WIDTH_CM
    chart.add_data(
        Reference(data, min_col=value_col, min_row=header_row, max_row=header_row + rows), titles_from_data=True
    )
    chart.set_categories(Reference(data, min_col=label_col, min_row=header_row + 1, max_row=header_row + rows))
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = True
    canvas.add_chart(chart, anchor)


def _block(ws: Worksheet, top: int, title: str, columns: list[tuple[str, str]], rows: list[dict[str, Any]]) -> int:
    """Write a titled table starting at row `top`, column A. Returns the header row."""
    ws.cell(row=top, column=1, value=title).font = BOLD
    header = top + 1
    for c, (name, _) in enumerate(columns, start=1):
        ws.cell(row=header, column=c, value=name).font = BOLD
    for r, row in enumerate(rows, start=header + 1):
        for c, (_, key) in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=c, value=_cell(row.get(key)))
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = "0%" if key in PERCENT_KEYS else "#,##0"
    return header


def _short(name: str | None, limit: int = 38) -> str:
    name = name or "-"
    return name if len(name) <= limit else name[: limit - 1] + "…"


def _charts_sheet(canvas: Worksheet, data: Worksheet, snap: Snapshot) -> None:
    """Charts stacked one per block on `canvas`; the tables they plot live on `data`."""
    anchors = (f"A{1 + i * CHART_ROW_STEP}" for i in range(100))

    modal = analysis.modal_summary(snap)
    head = _block(
        data,
        1,
        "Results by transport mode (annual line profit)",
        [
            ("Mode", "carrier"),
            ("Lines", "lines"),
            ("Vehicles", "vehicles"),
            ("Profit/year", "profitYear"),
            ("Profit per vehicle", "profitPerVehicle"),
            ("Loss-making lines", "losingLines"),
        ],
        modal,
    )
    n = len(modal)
    _bar_chart(canvas, data, "Annual profit by transport mode", 1, 4, n, next(anchors), head, horizontal=False)
    _bar_chart(canvas, data, "Profit per vehicle by transport mode", 1, 5, n, next(anchors), head, horizontal=False)
    _pie_chart(canvas, data, "Vehicles by transport mode", 1, 3, n, next(anchors), head)

    rows = [analysis.line_row(snap, ln) for ln in snap.lines]
    for r in rows:
        r["short"] = _short(r["name"])
    best = sorted(rows, key=lambda r: r["profitYear"] or 0, reverse=True)[:15]
    worst = sorted(rows, key=lambda r: r["profitYear"] or 0)[:15]
    cols = [
        ("Line", "short"),
        ("Mode", "carrier"),
        ("Vehicles", "vehicles"),
        ("Profit/year", "profitYear"),
        ("Current load", "load"),
    ]

    row = head + n + 3
    head = _block(data, row, "15 most profitable lines", cols, best)
    _bar_chart(canvas, data, "15 most profitable lines (profit/year)", 1, 4, len(best), next(anchors), head)

    row = head + len(best) + 3
    head = _block(data, row, "15 worst-performing lines", cols, worst)
    _bar_chart(canvas, data, "15 worst-performing lines (profit/year)", 1, 4, len(worst), next(anchors), head)

    row = head + len(worst) + 3
    dist = analysis.load_distribution(snap)
    head = _block(data, row, "Lines by load band (snapshot instant)", [("Band", "band"), ("Lines", "lines")], dist)
    _bar_chart(
        canvas, data, "Lines by load band (snapshot instant)", 1, 2, len(dist), next(anchors), head, horizontal=False
    )

    row = head + len(dist) + 3
    ops = [o for o in analysis.opportunities(snap, limit=100) if o["kind"] == "unshipped"][:12]
    for o in ops:
        o["label"] = _short(f"{o['producer']['name']} – {o['cargo']}", 45)
    head = _block(
        data,
        row,
        "Cargo produced but not shipped (per year)",
        [
            ("Producer – cargo", "label"),
            ("Produced/year", "producedYear"),
            ("Shipped/year", "shippedYear"),
            ("Unshipped/year", "unshippedYear"),
        ],
        ops,
    )
    _bar_chart(canvas, data, "Unshipped cargo per year", 1, 4, len(ops), next(anchors), head)

    data.column_dimensions["A"].width = 46
    for col in "BCDEF":
        data.column_dimensions[col].width = 18


def _apply_font(wb: Workbook) -> None:
    """Arial 12 everywhere; keeps bold where it was set."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.font = BOLD if cell.font is not None and cell.font.bold else REGULAR


def export_workbook(store: Store, snap: Snapshot, prev: Snapshot | None, folder: Path) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    clock = store.time_context(snap)

    ws = wb.active
    ws.title = "Summary"
    _summary(ws, snap, analysis.company_overview(snap, prev))
    for key, value in clock.time_basis.items():
        ws.append(["Time Basis: " + key, str(value)])
    ws.append(["Age unit", "GAME-YEAR = 365.25 calendar days; bracketed interpolation is an estimate"])

    charts = wb.create_sheet("Charts")
    _charts_sheet(charts, wb.create_sheet("Chart Data"), snap)

    lines = sorted(
        (analysis.line_row(snap, ln, clock=clock) for ln in snap.lines),
        key=lambda r: r["profitYear"] or 0,
        reverse=True,
    )
    ws = wb.create_sheet("Lines")
    _table(
        ws,
        [
            ("Line", "name"),
            ("Id", "id"),
            ("Mode", "carrier"),
            ("Stops", "stops"),
            ("Vehicles", "vehicles"),
            ("Profit/year", "profitYear"),
            ("Transported/year", "transportedYear"),
            ("Rate", "rate"),
            ("Current load", "load"),
            ("Issues", "issues"),
            ("Age (GAME-YEAR)", "ageYears"),
            ("Age evidence", "ageTime"),
        ],
        lines,
    )

    ws = wb.create_sheet("Vehicles")
    _table(
        ws,
        [
            ("Vehicle", "name"),
            ("Id", "id"),
            ("Mode", "carrier"),
            ("State", "state"),
            ("Line", "lineName"),
            ("Models", "models"),
            ("Age (GAME-YEAR)", "ageYears"),
            ("Age evidence", "ageTime"),
            ("Profit/year", "profitYear"),
            ("Value", "value"),
            ("Maintenance (0-1)", "maintenanceState"),
            ("No path", "noPath"),
            ("Stopped by player", "userStopped"),
            ("Days in depot", "daysInDepot"),
        ],
        sorted((analysis.vehicle_row(snap, v, clock=clock) for v in snap.vehicles), key=lambda r: r["profitYear"] or 0),
    )

    industry_rows = []
    for ind in snap.industries:
        row = analysis.industry_row(snap, ind)
        for direction, flows in (("input", row["inputs"]), ("output", row["outputs"])):
            for f in flows:
                industry_rows.append(
                    {
                        "name": row["name"],
                        "id": row["id"],
                        "level": row["level"],
                        "maxLevel": row["maxLevel"],
                        "closing": row["closing"],
                        "rating": row["productionRating"],
                        "direction": direction,
                        "cargo": f.get("cargo"),
                        "actual": f.get("producedYear") if direction == "output" else f.get("consumedYear"),
                        "max": f.get("maxProductionYear") if direction == "output" else f.get("maxConsumptionYear"),
                        "moved": f.get("shippedYear") if direction == "output" else f.get("deliveredYear"),
                    }
                )
    ws = wb.create_sheet("Industries")
    _table(
        ws,
        [
            ("Industry", "name"),
            ("Id", "id"),
            ("Level", "level"),
            ("Max level", "maxLevel"),
            ("Closing", "closing"),
            ("Production rating", "rating"),
            ("Flow", "direction"),
            ("Cargo", "cargo"),
            ("Produced/consumed per year", "actual"),
            ("Maximum per year", "max"),
            ("Shipped/delivered per year", "moved"),
        ],
        industry_rows,
    )

    ws = wb.create_sheet("Towns")
    _table(
        ws,
        [
            ("Town", "name"),
            ("Id", "id"),
            ("Population (capacity)", "population"),
            ("Residential/Commercial/Industrial", "capacities"),
            ("Line usage", "lineUsage"),
            ("Growing", "developmentActive"),
        ],
        sorted((analysis.town_row(t) for t in snap.towns), key=lambda r: r["population"] or 0, reverse=True),
    )

    ops = analysis.opportunities(snap, limit=100)
    ws = wb.create_sheet("Opportunities")
    kinds = {"unshipped": "Waiting cargo (produced but not shipped)", "needs_inputs": "Idle capacity (missing inputs)"}
    _table(
        ws,
        [
            ("Type", "kind"),
            ("Producer", "producer"),
            ("Cargo", "cargo"),
            ("Produced/year", "producedYear"),
            ("Capacity/year", "maxProductionYear"),
            ("Shipped/year", "shippedYear"),
            ("Unshipped/year", "unshippedYear"),
            ("Missing inputs", "missingInputs"),
            ("Nearest consumer", "consumer"),
            ("Distance (km)", "distanceKm"),
        ],
        [
            {
                **o,
                "kind": kinds[o["kind"]],
                "producer": o["producer"]["name"],
                "consumer": o["nearestConsumers"][0]["name"] if o["nearestConsumers"] else None,
                "distanceKm": o["nearestConsumers"][0]["distanceKm"] if o["nearestConsumers"] else None,
            }
            for o in ops
        ],
    )

    ws = wb.create_sheet("Findings")
    _table(
        ws,
        [
            ("Code", "code"),
            ("Severity", "severity"),
            ("Category", "category"),
            ("Title", "title"),
            ("Detail", "detail"),
            ("Suggestion", "suggestion"),
        ],
        [
            f.to_dict()
            for f in analysis.findings(snap, prev, history=store.recent_snapshots(24, snap.game_time), clock=clock)
        ],
    )
    history = store.company_history(limit=1000)
    ws = wb.create_sheet("History")
    _table(
        ws,
        [
            ("Date", "date_label"),
            ("Balance", "balance"),
            ("Loan", "loan"),
            ("Lines", "lines"),
            ("Vehicles", "vehicles"),
        ],
        history,
    )
    _line_chart(ws, "Balance and loan", "Money", [2, 3], len(history), "G2")
    _line_chart(ws, "Lines and vehicles", "Count", [4, 5], len(history), "G20")

    ws = wb.create_sheet("Decisions")
    _table(
        ws,
        [
            ("Id", "id"),
            ("Game date", "game_date"),
            ("Title", "title"),
            ("Rationale", "rationale"),
            ("Entities", "entities"),
            ("Status", "status"),
            ("Outcome", "outcome"),
        ],
        store.list_decisions(limit=1000),
    )

    _apply_font(wb)
    path = folder / f"tf3c_{snap.date_label}.xlsx"
    wb.save(path)
    return path
