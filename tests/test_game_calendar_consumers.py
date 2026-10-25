"""The same evidence-bound age across rows, facts, findings, workbook and history."""

from __future__ import annotations

import json
from contextlib import closing
from datetime import date

import pytest
from openpyxl import load_workbook

from conftest import base_snapshot
from tf3companion import analysis, reports
from tf3companion.export import export_workbook
from tf3companion.game_time import GameCalendarClock
from tf3companion.snapshot import Snapshot
from tf3companion.store import Store


def snapshot(tick, dated, schema=1, purchase=1000):
    raw = base_snapshot(game_time=tick)
    raw["meta"].update(schemaVersion=schema, yearDuration=91312)
    if dated is None:
        raw["meta"].pop("date", None)
    else:
        raw["meta"]["date"] = dict(zip(("year", "month", "day"), dated))
    raw["lines"] = [
        {
            "id": 1,
            "name": "Measured line",
            "vehicles": [2],
            "stops": [{}, {}],
            "profitYear": -100,
            "capacity": [],
            "issues": [],
        }
    ]
    raw["vehicles"] = [
        {
            "id": 2,
            "name": "Measured vehicle",
            "line": 1,
            "purchaseTime": purchase,
            "profitYear": -100,
            "carrier": "ROAD",
            "models": [],
        }
    ]
    return Snapshot(raw, source=f"fixture-{tick}")


@pytest.mark.parametrize("schema", [1, 2, 3])
def test_schema_compatibility_exact_three_years_and_no_evidence_mutation(tmp_path, schema):
    with closing(Store(tmp_path / "history.db")) as store:
        store.add_snapshot(snapshot(0, (1980, 1, 1), schema))
        latest = snapshot(4000, (1984, 1, 1), schema)
        store.add_snapshot(latest)
        store.add_decision("Observe", "No game action", snapshot=latest)
        before = [tuple(r) for r in store.conn.execute("SELECT * FROM snapshots ORDER BY game_time")]
        decisions = [tuple(r) for r in store.conn.execute("SELECT * FROM decisions")]
        clock = store.time_context(latest)
        line = analysis.line_row(latest, latest.lines[0], clock=clock)
        vehicle = analysis.vehicle_row(latest, latest.vehicles[0], clock=clock)
        assert line["ageYears"] == vehicle["ageYears"] == 3
        assert line["ageTime"]["value"] == 3
        assert line["timeBasis"]["anchorHash"] == vehicle["timeBasis"]["anchorHash"]
        assert latest.age_years(1000) is None
        assert latest.age_years(1000, clock=clock) == 3
        low = analysis.findings(latest, t=analysis.Thresholds(mature_years=4), clock=clock)
        high = analysis.findings(latest, t=analysis.Thresholds(mature_years=3), clock=clock)
        assert next(f for f in low if f.code == "line_losing_money").severity == "low"
        assert next(f for f in high if f.code == "line_losing_money").severity == "high"
        workbook = export_workbook(store, latest, None, tmp_path / "export")
        wb = load_workbook(workbook)
        for name in ("Lines", "Vehicles"):
            values = list(wb[name].values)
            index = values[0].index("Age (GAME-YEAR)")
            assert values[1][index] == 3
        assert clock.time_basis["anchorHash"] in str(list(wb["Summary"].values))
        assert [tuple(r) for r in store.conn.execute("SELECT * FROM snapshots ORDER BY game_time")] == before
        assert [tuple(r) for r in store.conn.execute("SELECT * FROM decisions")] == decisions


def test_old_purchase_unknown_and_replacement_does_not_inherit_observed_age(tmp_path):
    with closing(Store(tmp_path / "history.db")) as store:
        store.add_snapshot(snapshot(0, (1980, 1, 1), purchase=0))
        latest = snapshot(4000, (1984, 1, 1), purchase=3999)
        store.add_snapshot(latest)
        clock = store.time_context(latest)
        assert analysis.line_age_years(latest, latest.lines[0], clock=clock) == pytest.approx(0.001)
        latest.raw["vehicles"][0]["purchaseTime"] = -1
        row = analysis.line_row(latest, latest.lines[0], clock=clock)
        assert row["ageYears"] is None and row["ageTime"]["reason"] == "outside-observed-calendar"
        loss = next(f for f in analysis.findings(latest, clock=clock) if f.code == "line_losing_money")
        assert loss.severity == "low"
        assert "younger" not in loss.detail and "unknown" in loss.detail


def test_historical_context_excludes_future_and_detects_ingestion_rewind(tmp_path):
    with closing(Store(tmp_path / "history.db")) as store:
        first = snapshot(0, (1980, 1, 1))
        middle = snapshot(2000, (1982, 1, 1))
        latest = snapshot(4000, (1984, 1, 1))
        for s in (first, middle, latest):
            store.add_snapshot(s)
        clock = store.time_context(middle)
        assert [t for t, _ in clock.anchors] == [0, 2000]
        assert clock.calendar_at_tick(3000).reason == "future-timestamp"
        store.add_snapshot(middle)
        assert store.time_context(latest).elapsed_between_ticks(1000, 4000).reason == "ambiguous-timeline"


def test_calendar_baseline_not_nominal_tick_subtraction(tmp_path):
    with closing(Store(tmp_path / "history.db")) as store:
        for s in (snapshot(0, (1980, 1, 1)), snapshot(3000, (1983, 1, 1)), snapshot(4000, (1984, 1, 1))):
            store.add_snapshot(s)
        latest = store.latest()
        # 365 days in 1983 is short of the 365.25-day target; 1980 is the latest supported baseline.
        baseline = store.calendar_baseline(latest)
        assert baseline.game_time == 0
        assert latest.year_duration == 91312
        md = reports.render_report(store, "decisions")
        assert "window complete" not in md


def test_store_refuses_conversion_across_save_identities(tmp_path):
    with closing(Store(tmp_path / "history.db")) as store:
        store.add_snapshot(snapshot(0, (1980, 1, 1)))
        store.add_snapshot(snapshot(4000, (1984, 1, 1)))
        latest = store.latest()
        assert store.time_context(latest).elapsed_between_ticks(0, 4000).value == pytest.approx(4, abs=0.01)
        # Legacy rows can carry a different identity; the store API refuses to write them, so alter evidence in place.
        store.conn.execute(
            "UPDATE snapshots SET data = json_set(data, '$.meta.saveId', 'other-save') WHERE game_time = 0"
        )
        store.conn.execute(
            "UPDATE snapshots SET data = json_set(data, '$.meta.saveId', 'this-save') WHERE game_time = 4000"
        )
        store.conn.commit()
        latest = store.latest()
        clock = store.time_context(latest)
        assert clock.ambiguous
        measure = clock.elapsed_between_ticks(0, 4000)
        assert measure.value is None and measure.reason == "ambiguous-timeline"
        assert "Ambiguous timeline" in " ".join(clock.time_basis["warnings"])
        assert store.calendar_baseline(latest) is None


def test_store_refuses_conversion_across_calendar_reversal(tmp_path):
    with closing(Store(tmp_path / "history.db")) as store:
        store.add_snapshot(snapshot(0, (1980, 3, 1)))
        store.add_snapshot(snapshot(2000, (1980, 1, 1)))
        latest = store.latest()
        clock = store.time_context(latest)
        assert clock.ambiguous
        assert clock.elapsed_between_ticks(0, 2000).reason == "ambiguous-timeline"
        assert store.calendar_baseline(latest) is None
        assert latest.age_years(0, clock=clock) is None


def test_store_refuses_conversion_when_a_tick_is_reingested_with_another_date(tmp_path):
    with closing(Store(tmp_path / "history.db")) as store:
        store.add_snapshot(snapshot(0, (1980, 1, 1)))
        store.add_snapshot(snapshot(2000, (1982, 1, 1)))
        store.add_snapshot(snapshot(4000, (1984, 1, 1)))
        before = [tuple(r) for r in store.conn.execute("SELECT * FROM snapshots ORDER BY game_time")]
        assert store.snapshot_count() == 3
        # One row per tick: the same tick can never hold two dates, so a re-ingest with
        # another date replaces it; being older than the newest tick it is a rewind.
        store.add_snapshot(snapshot(2000, (1990, 1, 1)))
        assert store.snapshot_count() == 3
        clock = store.time_context(store.latest())
        assert clock.ambiguous
        assert clock.elapsed_between_ticks(0, 4000).reason == "ambiguous-timeline"
        assert before != [tuple(r) for r in store.conn.execute("SELECT * FROM snapshots ORDER BY game_time")]


def test_time_context_preserves_raw_json_and_decisions(tmp_path):
    with closing(Store(tmp_path / "history.db")) as store:
        first = snapshot(0, (1980, 2, 1))
        store.add_snapshot(first)
        second = snapshot(1000, (1980, 1, 1))
        store.add_snapshot(second)
        store.add_decision("Observe", "No game action", snapshot=second)
        rows = [tuple(r) for r in store.conn.execute("SELECT * FROM snapshots ORDER BY game_time")]
        decisions = [tuple(r) for r in store.conn.execute("SELECT * FROM decisions")]
        for s in (first, second):
            store.time_context(s)
        assert [tuple(r) for r in store.conn.execute("SELECT * FROM snapshots ORDER BY game_time")] == rows
        assert [tuple(r) for r in store.conn.execute("SELECT * FROM decisions")] == decisions
        assert store.latest().raw["meta"]["date"] == {"year": 1980, "month": 1, "day": 1}


def test_age_years_supported_zero_unknown_past_and_future(tmp_path):
    with closing(Store(tmp_path / "history.db")) as store:
        store.add_snapshot(snapshot(1000, (1980, 1, 1)))
        latest = snapshot(5000, (1984, 1, 1))
        store.add_snapshot(latest)
        clock = store.time_context(latest)
        assert latest.age_years(5000, clock=clock) == 0  # supported zero, not unknown
        assert latest.age_years(3000, clock=clock) == pytest.approx(
            (date(1984, 1, 1).toordinal() - date(1982, 1, 1).toordinal()) / 365.25, abs=0.01
        )
        assert latest.age_years(1000, clock=clock) == pytest.approx(4, abs=0.01)
        assert latest.age_years(999, clock=clock) is None  # before the first anchor
        assert latest.age_years(5001, clock=clock) is None  # purchased after the snapshot
        assert latest.age_years(None, clock=clock) is None
        assert latest.age_years(3000) is None  # no implicit nominal clock
        undated = snapshot(9000, None)
        assert undated.age_years(9000, clock=GameCalendarClock.from_samples([(9000, None)])) is None


def _surface_snapshot(tick, dated):
    raw = snapshot(tick, dated).raw
    raw["lines"] = [
        {
            "id": i,
            "name": f"Line {i}",
            "vehicles": [10 + i],
            "stops": [{"stationGroup": 500}, {"stationGroup": 501}],
            "profitYear": -10 * i,
            "capacity": [],
            "issues": [],
        }
        for i in (1, 2, 3)
    ]
    raw["vehicles"] = [
        {
            "id": 10 + i,
            "name": f"Vehicle {i}",
            "line": i,
            "purchaseTime": purchase,
            "profitYear": -10 * i,
            "carrier": "ROAD",
            "models": [],
        }
        for i, purchase in ((1, 0), (2, 2999), (3, -1))
    ]
    return raw


@pytest.fixture
def surfaces(tmp_path, monkeypatch):
    from conftest import write_snapshot
    from tf3companion import mcp_server
    from tf3companion.service import Companion

    data, userdata = tmp_path / "data", tmp_path / "userdata"
    monkeypatch.setenv("TF3C_DATA_DIR", str(data))
    monkeypatch.setenv("TF3C_USERDATA_DIR", str(userdata))
    folder = userdata / "towns_industries"
    write_snapshot(folder, _surface_snapshot(0, (1980, 1, 1)))
    write_snapshot(folder, _surface_snapshot(4000, (1984, 1, 1)))
    companion = Companion()
    monkeypatch.setattr(mcp_server, "open_companion", lambda save=None: Companion(save=save))
    yield companion, mcp_server
    companion.close()


def test_same_unrounded_age_and_time_basis_on_every_surface(surfaces, capsys):
    from tf3companion.cli import main

    companion, mcp = surfaces
    ages = {r["id"]: r["ageYears"] for r in companion.lines("ageYears", 10)}
    assert ages[1] == pytest.approx(4, abs=0.01) and ages[3] is None
    assert 0 < ages[2] < ages[1] and ages[2] != round(ages[2], 2)
    anchor_hash = companion.overview()["timeBasis"]["anchorHash"]

    def keyed(rows):
        return {r["id"]: (r["ageYears"], r["ageTime"]["value"], r["timeBasis"]["anchorHash"]) for r in rows}

    expected = {i: (ages[i], ages[i], anchor_hash) for i in ages}
    assert keyed(companion.lines("ageYears", 10)) == expected
    assert keyed(mcp.list_lines("ageYears", 10)) == expected
    assert main(["lines", "--sort", "ageYears", "--limit", "10"]) == 0
    cli_rows = json.loads(capsys.readouterr().out)
    assert keyed(cli_rows) == expected
    # Detail surfaces share the same measurement for the line and its vehicle rows.
    for detail in (companion.line_detail(1), mcp.line_detail("1")):
        assert detail["ageYears"] == ages[1] and detail["timeBasis"]["anchorHash"] == anchor_hash
        assert detail["vehicleList"][0]["ageYears"] == ages[1]
    assert main(["line", "1"]) == 0
    cli_detail = json.loads(capsys.readouterr().out)
    assert cli_detail["ageYears"] == ages[1] and cli_detail["timeBasis"]["anchorHash"] == anchor_hash
    assert keyed(mcp.list_vehicles("all", None, 10))[11][0] == ages[1]


def test_age_sorting_places_unknown_ages_last_on_every_surface(surfaces, capsys):
    from tf3companion.cli import main

    companion, mcp = surfaces
    assert [r["id"] for r in companion.lines("ageYears", 10)] == [1, 2, 3]
    assert [r["id"] for r in mcp.list_lines("ageYears", 10)] == [1, 2, 3]
    assert main(["lines", "--sort", "ageYears", "--limit", "10"]) == 0
    assert [r["id"] for r in json.loads(capsys.readouterr().out)] == [1, 2, 3]
    assert [r["id"] for r in companion.lines("ageYears", 2)] == [1, 2]


def test_workbook_leaves_unknown_age_cells_blank(surfaces, tmp_path):
    companion, _ = surfaces
    snap = companion.current()
    workbook = load_workbook(export_workbook(companion.store, snap, None, tmp_path / "export"))
    for name, id_column in (("Lines", "Id"), ("Vehicles", "Id")):
        values = list(workbook[name].values)
        header = values[0]
        by_id = {row[header.index(id_column)]: row[header.index("Age (GAME-YEAR)")] for row in values[1:]}
        known = {k: v for k, v in by_id.items() if v is not None}
        assert len(known) == 2 and len(by_id) == 3
        assert all(isinstance(v, (int, float)) for v in known.values())
        evidence = {row[header.index(id_column)]: row[header.index("Age evidence")] for row in values[1:]}
        unknown = next(k for k, v in by_id.items() if v is None)
        assert "outside-observed-calendar" in str(evidence[unknown])
